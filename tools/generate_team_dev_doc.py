from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "牧融绿链_组员开发说明文档_公网功能融合版.docx"


def set_font(run, size=10.5, bold=False, color=None):
    run.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_font(
        run,
        size=16 if level == 1 else 13,
        bold=True,
        color=RGBColor(31, 78, 121) if level == 1 else RGBColor(47, 84, 150),
    )
    return paragraph


def add_para(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.24)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run("• " + item)
        set_font(run)


def cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    set_font(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell_text(cell, header, bold=True)
        shade_cell(cell, "D9EAF7")
        if widths:
            cell.width = Inches(widths[index])
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph()
    return table


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    doc.styles["Normal"].font.name = "微软雅黑"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("牧融绿链组员开发说明文档")
    set_font(run, size=22, bold=True, color=RGBColor(31, 78, 121))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("给组员看的项目目标、架构、技术栈和后续分工说明")
    set_font(run, size=12, color=RGBColor(89, 89, 89))

    doc.add_paragraph()
    add_para(doc, "本文档面向项目组内部成员，目的是帮助大家快速理解牧融绿链项目的目标、技术栈、代码结构、数据流、页面模块、接口和后续分工。它不是答辩稿，不强调展示话术，而强调项目怎么开发、怎么维护、怎么继续补功能。")

    add_heading(doc, "一、项目目标")
    add_para(doc, "牧融绿链是一个面向“工行杯”比赛场景的高原畜牧绿色金融风险评估、产业链资金闭环与贷后管理平台。系统必须突出工行作为绿色金融服务主体的角色，而不是做成泛泛的高原畜牧展示平台。")
    add_para(doc, "一句话目标：融合气象遥感、经营主体、工行授信、产业链订单、资金流向、还款逾期和保险协同数据，为工行绿色信贷提供可解释的风险评估、产业链资金闭环、贷后预警和处置建议。")
    add_bullets(
        doc,
        [
            "服务对象：工行客户经理、风控人员、绿色金融业务人员。",
            "核心业务：授信准入、额度判断、产业链资金追踪、贷后监测、风险处置、银保协同、绿色金融绩效评价。",
            "系统输出：风险分、风险等级、主风险因子、证据链、处置动作、数据缺口提示。",
            "开发底线：所有 sample/demo/弱标签都必须明确标注，不能伪造成真实生产数据。",
        ],
    )

    add_heading(doc, "二、技术栈")
    add_table(
        doc,
        ["层级", "技术", "说明"],
        [
            ["前端", "Vue 3 + ECharts + 原生 CSS", "无构建工具，直接通过静态 HTML/JS/CSS 运行，适合比赛演示。"],
            ["后端", "Python FastAPI", "提供 API、静态页面服务、文件上传、数据质量、模型接口。"],
            ["模型", "规则模型 + sklearn", "小样本用规则模型，样本足够时使用 RandomForest + Ridge 混合模型。"],
            ["数据", "JSON Store", "主要数据存放在 backend/data_store/*.json，方便本地演示和迁移。"],
            ["数据库预留", "MySQL", "后端 db.py、schema.sql 已预留，当前默认回退 JSON Store。"],
            ["公开数据处理", "Python 脚本", "public_data/scripts/ingest_public_data.py 用于处理 TPDC/MODIS/CMFD 等数据。"],
        ],
        widths=[1.2, 1.8, 4.1],
    )

    add_heading(doc, "三、目录结构")
    add_table(
        doc,
        ["路径", "作用"],
        [
            ["backend/server.py", "FastAPI 后端入口，定义 API、静态页面服务、上传、外部 API 配置。"],
            ["backend/data.py", "业务数据访问层，负责平台核心数据、模块、品牌定位和基础风险评估。"],
            ["backend/models.py", "模型层，负责训练、预测、特征重要性、趋势预测和标签说明。"],
            ["backend/store.py", "JSON Store 读写和 CSV 导入，负责数据质量来源追踪。"],
            ["backend/data_store/*.json", "当前系统数据文件，包括气象、遥感、主体、授信、标签等。"],
            ["frontend/index.html", "主前台 SPA 页面结构。"],
            ["frontend/app.js", "Vue 逻辑、风险评估工作台、图表、模型状态、页面交互。"],
            ["frontend/styles.css", "前台样式，包含首页、工作台、数据底座、路线图等。"],
            ["frontend/admin.html", "管理端页面，负责 CSV 上传、字段校验、图片轮播等。"],
            ["public_data/processed", "已处理 CSV 数据。"],
            ["public_data/raw", "原始公开数据。"],
            ["public_data/samples", "样例 CSV，可用于导入演示。"],
            ["public_data/scripts", "公开数据转换和样例生成脚本。"],
        ],
        widths=[2.2, 4.8],
    )

    add_heading(doc, "四、启动方式")
    add_para(doc, "进入项目目录：")
    add_para(doc, "cd D:\\工行杯\\yak-risk-platform")
    add_para(doc, "启动后端：")
    add_para(doc, "python -m uvicorn backend.server:app --host 127.0.0.1 --port 8000")
    add_para(doc, "常用访问地址：")
    add_bullets(
        doc,
        [
            "前台：http://127.0.0.1:8000",
            "风险评估：http://127.0.0.1:8000/#dashboard",
            "管理端：http://127.0.0.1:8000/admin",
            "API 文档：http://127.0.0.1:8000/docs",
        ],
    )
    add_para(doc, "注意：run.err.log 里如果出现 MySQL root 无密码导致的连接失败，这是正常现象。系统会自动回退到 JSON Store，不影响演示。")

    add_heading(doc, "五、当前数据状态")
    add_table(
        doc,
        ["数据表", "当前规模", "说明"],
        [
            ["weather_data.json", "1500 行", "TPDC CMFD，25 县 × 2020-2024 共 60 个月。"],
            ["remote_sensing_data.json", "1500 行", "MODIS/TPDC 遥感，含 NDVI、积雪、退化、载畜量字段。"],
            ["business_subjects.json", "75 行", "25 县 × 3 类样例主体，sample/demo。"],
            ["finance_credit.json", "75 行", "工行授信、用信、还款、保险样例，sample/demo。"],
            ["forage_supply_demand.json", "126 行", "Geodoi 宏观饲草供需，只做背景，不参与训练。"],
            ["risk_event_labels.json", "25 行", "demo 风险事件标签，只用于跑通标签链路。"],
            ["supply_chain_orders", "接口预留", "产业链订单表，公网展示字段包含 order_id、subject_name、order_type、supplier、amount、linked_credit_id。"],
            ["supply_chain_payments", "接口预留", "产业链资金流向表，公网展示字段包含 payment_id、order_id、from_account、to_account、amount、channel、status。"],
        ],
        widths=[2.0, 1.1, 3.9],
    )
    add_para(doc, "组员注意：真实环境数据可以用于模型训练；经营主体、授信台账和风险标签目前是样例，任何页面或文案都不能宣称它们是真实银行生产数据。")

    add_heading(doc, "六、核心数据流")
    add_para(doc, "系统的数据闭环如下：")
    add_bullets(
        doc,
        [
            "管理端上传 CSV 或直接修改 backend/data_store/*.json。",
            "store.py 负责解析 CSV、字段校验、写入 JSON Store 和记录导入元数据。",
            "data.py 读取最新数据，合成平台需要的 regions、subjects、finance、risk_assessment 等结构。",
            "models.py 从气象、遥感、经营、金融四类数据中抽取特征，按 region_id + month 训练模型。",
            "server.py 将平台数据和模型结果通过 /api/platform、/api/model/* 等接口返回给前端。",
            "frontend/app.js 根据接口数据渲染对象池、风险报告、证据链、图表和处置建议。",
            "产业链、绿色绩效、边疆民生等页面当前以业务结构和接口预留为主，后续需要把真实订单、资金流水、碳账户、GEP、保险协同数据接进来。",
        ],
    )

    add_heading(doc, "七、模型逻辑")
    add_para(doc, "模型训练单位是 region_id + month。当前共有 25 个县、60 个月，形成 1500 个县域月度样本。")
    add_table(
        doc,
        ["维度", "特征示例"],
        [
            ["气象", "温度、降水、风速、雪深、寒潮风险、暴雪风险、干旱风险。"],
            ["遥感", "NDVI、NDVI 变化、植被覆盖、积雪覆盖、退化等级、载畜量。"],
            ["经营", "主体评分、保险覆盖率。"],
            ["金融", "还款状态、逾期次数、用信风险。"],
        ],
        widths=[1.2, 5.8],
    )
    add_para(doc, "当前模型状态：ml_hybrid，样本数 1500，特征数 16，label_type=rule_label。也就是说，现在是“真实环境数据 + 规则弱标签”的评分模型，不是真实灾害预测或真实逾期预测。")

    add_heading(doc, "八、主要 API")
    add_table(
        doc,
        ["接口", "作用"],
        [
            ["/api/health", "健康检查。"],
            ["/api/platform", "前端主接口，返回平台全量数据。"],
            ["/api/model/status", "查看模型状态、样本数、真实数据比例、警告信息。"],
            ["/api/model/train", "重新训练模型。"],
            ["/api/model/predict", "返回所有区域预测结果。"],
            ["/api/model/importance", "返回特征重要性。"],
            ["/api/model/forecast", "返回某个区域未来风险趋势。"],
            ["/api/data-quality", "返回各数据表质量报告。"],
            ["/api/import/csv", "管理端 CSV 导入接口。"],
            ["/api/public-data/samples", "返回公开样例 CSV 列表。"],
            ["/api/integrations/status", "查看外部 API 配置状态和可接入能力。"],
        ],
        widths=[2.2, 4.8],
    )

    add_heading(doc, "九、前端页面怎么理解")
    add_table(
        doc,
        ["页面", "职责"],
        [
            ["首页", "介绍项目定位、工行业务闭环和核心能力。"],
            ["平台概览", "展示接入数据量、覆盖县域、监测时间、风险分布等总览指标。"],
            ["业务模块", "展示授信准入、生态监测、贷后处置、银保协同、产业链追踪等模块。"],
            ["数据底座", "展示气象、遥感、主体、授信、标签等数据资产和质量。"],
            ["实施路线", "展示项目落地路径和后续建设阶段。"],
            ["风险评估 /#dashboard", "核心页面，服务客户经理做风险评估和处置。"],
            ["保险协同", "围绕工行证据链、保险覆盖、理赔联动和风险分担设计。"],
            ["产业链", "展示产业链资金闭环：工行放款、饲草采购、物流配送、牧场入库、产品回款、贷后核销。"],
            ["绿色绩效", "展示 NDVI 改善、县域风险分布、绿色信贷指标、碳账户/GEP 接口预留。"],
            ["边疆民生", "展示边境重点县、风险状态、政府/工行/保险/合作社四方协同。"],
            ["管理端 /admin", "CSV 导入、数据管理、字段校验、轮播图管理。"],
        ],
        widths=[1.8, 5.2],
    )

    add_heading(doc, "十、页面模块开发要点")
    add_table(
        doc,
        ["模块", "开发要点"],
        [
            ["风险评估", "保持风控终端风格，围绕对象池、风险报告、证据链、处置动作和评估明细组织页面。"],
            ["产业链", "先把 supply_chain_orders 和 supply_chain_payments 的 CSV 导入、字段校验、列表展示做通，再接资金闭环图。"],
            ["保险协同", "补充保单号、承保机构、保障额度、理赔状态、保险到期日，和风险处置建议联动。"],
            ["绿色绩效", "不要只写口号，要围绕 NDVI 趋势、绿色信贷余额、支持牧户数、减灾减损、碳减排当量、GEP 接口预留组织。"],
            ["边疆民生", "重点展示边境重点县、平均风险、重点监测名单和政府/银行/保险/合作社协同分工。"],
            ["数据底座", "做成数据资产目录，突出数据来源、字段、规模、质量、用途、sample/demo 标识。"],
            ["管理端", "优化 CSV 上传体验，支持导入预览、缺字段提示、错误行定位、导入记录追踪。"],
        ],
        widths=[1.6, 5.4],
    )

    add_heading(doc, "十一、开发分工建议")
    add_table(
        doc,
        ["角色", "建议任务"],
        [
            ["前端同学", "维护 index.html、app.js、styles.css；优化风险工作台、主体详情、产业链、绿色绩效、边疆民生和响应式样式。"],
            ["后端同学", "维护 server.py、store.py；补充 API、CSV 导入校验、数据质量报告、产业链数据表和异常处理。"],
            ["算法同学", "维护 models.py；改进特征工程、模型评估、标签说明和风险解释。"],
            ["数据同学", "维护 public_data 和 data_store；整理 TPDC、MODIS、CMFD、Geodoi、产业链订单、保险协同数据来源和字段映射。"],
            ["答辩/文档同学", "整理演示流程、项目说明、数据边界、创新点和常见问答。"],
        ],
        widths=[1.5, 5.5],
    )

    add_heading(doc, "十二、开发注意事项")
    add_bullets(
        doc,
        [
            "不要把系统做成政府汇报驾驶舱或 PPT 展示页，要始终突出工行客户经理的风控动作和资金闭环动作。",
            "不要把 sample/demo/弱标签写成真实生产数据，尤其是经营主体、授信台账和风险标签。",
            "前端改动优先保持现有 Vue 3 CDN 结构，不要临时引入复杂构建工具。",
            "后端优先保持 JSON Store 可用，MySQL 不可用时必须能正常回退。",
            "模型说明必须保留 label_type=rule_label 的边界，不要宣称真实逾期预测已经完成。",
            "产业链、碳账户、GEP、保险理赔等暂未真实接入的数据，页面必须写清楚接口预留或 sample/demo。",
            "每次改 app.js 后运行 node --check frontend\\app.js。",
            "每次改 Python 后运行 ast 语法检查，确保 server.py、data.py、store.py、models.py 都能解析。",
        ],
    )

    add_heading(doc, "十三、验证命令")
    add_para(doc, "前端语法检查：")
    add_para(doc, "node --check frontend\\app.js")
    add_para(doc, "Python 语法检查：")
    add_para(doc, "python -c \"import ast, pathlib; [ast.parse(pathlib.Path(f).read_text(encoding='utf-8')) for f in ['backend/server.py','backend/data.py','backend/store.py','backend/models.py']]; print('python syntax ok')\"")
    add_para(doc, "接口检查：")
    add_bullets(
        doc,
        [
            "Invoke-WebRequest -UseBasicParsing http://127.0.0.1:8000/api/health",
            "Invoke-WebRequest -UseBasicParsing http://127.0.0.1:8000/api/platform",
            "Invoke-WebRequest -UseBasicParsing http://127.0.0.1:8000/admin",
        ],
    )

    add_heading(doc, "十四、后续开发优先级")
    add_bullets(
        doc,
        [
            "补充真实灾害、理赔、逾期标签，逐步从规则弱标签走向真实监督模型。",
            "完善客户经理处置记录，让贷后核查结果可以回流到模型样本。",
            "补齐 supply_chain_orders 和 supply_chain_payments 两类产业链数据，形成可导入、可展示、可追踪的资金闭环。",
            "把绿色绩效中的绿色信贷余额、支持牧户数、碳减排当量、GEP、可持续评分从接口预留逐步接到真实或脱敏数据。",
            "完善保险协同数据，包括承保额度、保单期限、理赔状态、银保联动处置记录。",
            "增强管理端 CSV 上传体验，包括字段映射、错误定位、导入预览。",
            "把数据底座页面改得更像数据资产目录，突出来源、字段、质量和用途。",
            "如果有稳定服务器，再恢复公网部署；答辩优先使用本地演示，避免 VPS 失联风险。",
        ],
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("牧融绿链 - 组员开发说明文档")
        set_font(run, size=9, color=RGBColor(128, 128, 128))

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
